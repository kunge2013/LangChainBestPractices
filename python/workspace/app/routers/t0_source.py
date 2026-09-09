"""
Router for T0 source data import (2.5).
"""
from __future__ import annotations

import io

from fastapi import APIRouter, Depends, File, UploadFile
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.database import get_db
from app.schemas.common import ApiResponse
from app.services.t0_service import import_source_data

router = APIRouter(prefix="/api/v1/ledger-t0-source", tags=["T0基础数据导入"])


@router.post("/import", response_model=ApiResponse)
def import_t0_source(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
):
    """Import T0 source data from Excel file."""
    content = file.file.read()
    wb = load_workbook(io.BytesIO(content), data_only=True)

    # Look for sheet named '未清明细'
    sheet = None
    for name in wb.sheetnames:
        if "未清明细" in name:
            sheet = wb[name]
            break
    if sheet is None:
        sheet = wb.active

    rows = []
    headers = None
    for row in sheet.iter_rows(values_only=True):
        if headers is None:
            headers = [str(c).strip() if c else "" for c in row]
            continue
        row_dict = {}
        for i, val in enumerate(row):
            if i < len(headers):
                row_dict[headers[i]] = val
        rows.append(row_dict)

    count = import_source_data(db, rows)
    db.commit()

    return ApiResponse(
        code="0",
        message=f"导入成功，共导入 {count} 条记录",
        data={"count": count},
    )
